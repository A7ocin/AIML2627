"""Rebuild import workbooks and instructor guides from question-bank.json.

python -m pip install openpyxl beautifulsoup4
python tools/build.py [--course-root /path/to/AIML]
"""
from pathlib import Path
import argparse,hashlib,html,json,collections,math
from bs4 import BeautifulSoup as BS
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

HERE=Path(__file__).resolve().parents[1]
parser=argparse.ArgumentParser();parser.add_argument('--course-root',type=Path,default=HERE.parent);args=parser.parse_args()
COURSE=args.course_root.resolve()
bank=json.loads((HERE/'question-bank.json').read_text(encoding='utf-8'))
fmt=json.loads((HERE/'tools/import-format.json').read_text(encoding='utf-8'))
questions=bank['questions'];sources={};errors=[]
H=html.escape

for lecture in bank['lectures']:
    pack=lecture['pack'];p=COURSE/pack/'index.html';raw=p.read_bytes()
    slides=BS(raw.decode('utf-8'),'html.parser').select('.slides>section')
    handout=BS((COURSE/pack/'handout.html').read_text(encoding='utf-8'),'html.parser')
    ids={s.get('data-slide',s.get('id')):(i+1,s) for i,s in enumerate(slides)}
    sources[pack]={'slides_sha256':hashlib.sha256(raw).hexdigest(),'handout_sha256':hashlib.sha256((COURSE/pack/'handout.html').read_bytes()).hexdigest(),'slides':len(slides)}
    selected=[q for q in questions if q['lecture']==lecture['number']]
    assert len({q['id'] for q in selected})==len(selected)
    assert selected, lecture['pack']
    for q in selected:
        assert q['slide'] in ids,(q['id'],q['slide'])
        n,s=ids[q['slide']];q['slide_number']=n
        head=s.select_one('h2,h3,h1');q['slide_title']=head.get_text(' ',strip=True) if head else q['slide']
        q['slide_href']=f'../{pack}/index.html#/{n-1}'
        chapter=handout.select_one('[data-source-slide="'+q['slide']+'"]');assert chapter,q['id']
        q['handout_href']=f'../{pack}/handout.html#{chapter["id"]}'
        assert q['when'] in ['before','after']
        assert q['type'] in fmt['types_used']
        assert 45<=q['seconds']<=240
        assert q['stem'].strip() and q['explanation'].strip() and q['misconception'].strip()
        assert len(q['stem'])<=600,(q['id'],'long stem')
        assert len(q['choices'])==(4 if q['type']=='MCQ' else 0)
        if q['type']=='MCQ':
            assert len(set(q['choices']))==4,q['id']
            assert q['choices'][q['correct_index']-1]==q['correct'],q['id']
            assert all(len(c)<260 for c in q['choices']),q['id']
        else:assert q['correct'] is None and q['correct_index'] is None
    core=sorted([q for q in selected if q['core']],key=lambda q:q['id'])
    ordering=[(q['slide_number'],q['when']=='after') for q in core]
    assert ordering==sorted(ordering),(pack,'core teaching order',ordering)

assert len(questions)==92 and sum(q['core'] for q in questions)==68
assert len({q['stem'] for q in questions})==len(questions)

def export(path,items):
    path.parent.mkdir(exist_ok=True,parents=True)
    wb=Workbook();ws=wb.active;ws.title='Questions'
    ws.append(fmt['columns'])
    for q in items:
        row=[q['type'],q['id']+' · '+q['stem'],str(q['correct_index']) if q['correct_index'] else None]+q['choices']
        ws.append(row+[None]*(10-len(row)))
    ws.freeze_panes='D2';ws.auto_filter.ref=f'A1:J{ws.max_row}'
    for cell in ws[1]:
        cell.font=Font(name='Calibri',size=11,bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='163345');cell.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=26
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='Calibri',size=11,color='243549');cell.alignment=Alignment(wrap_text=True,vertical='top')
            if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='EEF7F4')
        ws.row_dimensions[row[0].row].height=105
    for i,width in enumerate([19,86,12]+[40]*7,1):ws.column_dimensions[get_column_letter(i)].width=width
    wb.properties.title=path.stem+' · Wooclap question import';wb.properties.creator='AIML course'
    wb.save(path)
    # Verify the saved workbook, not only the authoring data.
    check=load_workbook(path,data_only=False)
    assert check.sheetnames==['Questions']
    got=list(check.active.iter_rows(values_only=True))
    assert list(got[0])==fmt['columns']
    assert len(got)==len(items)+1
    for actual,q in zip(got[1:],items):
        assert actual[0]==q['type'] and actual[1]==q['id']+' · '+q['stem']
        assert actual[2]==(str(q['correct_index']) if q['correct_index'] else None)
        assert list(actual[3:3+len(q['choices'])])==q['choices']
        assert all(x is None for x in actual[3+len(q['choices']):])
    assert not any(c.data_type=='f' for row in check.active for c in row)
    check.close()

rows=[];content=[];md=[]
intro='''These are instructor materials with answer keys. The 68 core questions are distributed across sixteen lectures; the 24 optional questions are separate imports. All questions are in English. Each MCQ has one correct answer; open responses are ungraded and have a suggested rubric here. Times include voting and a brief debrief, and are planning estimates, not automatic Wooclap timer settings.'''
method='''For a “before” question, collect responses before showing the explanatory slide or running the demonstration. For an “after” question, check whether the explanation transfers to a concrete case. Ask students to commit individually, optionally discuss a split vote in pairs, then explain the reasoning. This is a first undergraduate introduction. Before-topic questions use everyday examples or ideas explicitly explained earlier; invite guesses and explain the answer after voting. Optional questions are basic extensions, not prerequisites or trick questions. Treat opening questions as ungraded diagnostics, not presumed prior mastery. Skip an optional item when the core question already provides sufficient evidence.'''
importnote='''In a Wooclap event, choose Import questions → Import from an Excel file and select that lecture’s core workbook. Import the optional workbook only if you want those additional items. Question IDs connect the imported titles to this guide; the slide placement and timing are teaching instructions, not fields imported by Wooclap. Open questions have blank Correct cells intentionally. After importing, preview the answer key, mathematical characters and question order, and check the question settings. The existing course events have already been imported. For those events, update questions in place to retain embed links; importing again would add duplicate questions. Revised live content is checked separately in live-revision-validation.json.'''
formatnote='''The official help page documents template-based Excel import. The workbooks use its published Type / Title / Correct / Choice schema, checked against an archived template copy provided by Université de Poitiers (June 2024). The current help page does not expose the template download without entering the event workflow. Only the two verified basic types, MCQ and OpenQuestion, are used. Single-answer MCQ keys use one-based choice positions. The AI Importer is not required for these files.'''
for lecture in bank['lectures']:
    qs=sorted([q for q in questions if q['lecture']==lecture['number']],key=lambda q:q['id'])
    core=[q for q in qs if q['core']];opt=[q for q in qs if not q['core']]
    stem=lecture['pack'];export(HERE/'imports'/f'{stem}.xlsx',core);export(HERE/'imports/optional'/f'{stem}_optional.xlsx',opt)
    minutes=sum(q['seconds'] for q in core)/60
    rows.append(f'<tr><td><a href="#lecture-{lecture["number"]}">{lecture["number"]:02d} · {H(lecture["title"])}</a></td><td>{len(core)}</td><td>{len(opt)}</td><td>{minutes:g} min</td><td><a href="imports/{stem}.xlsx">Core Excel</a> · <a href="imports/optional/{stem}_optional.xlsx">Optional Excel</a></td></tr>')
    block=[f'<section class="lecture" id="lecture-{lecture["number"]}"><h2>{lecture["number"]:02d} · {H(lecture["title"])}</h2><p class="lecture-meta">{len(core)} core + {len(opt)} optional · Core discussion budget: approximately {minutes:g} minutes</p>']
    md.extend([f'## {lecture["number"]:02d} · {lecture["title"]}','',f'{len(core)} core + {len(opt)} optional. Core budget: approximately {minutes:g} minutes.','',f'[Core import](imports/{stem}.xlsx) · [Optional import](imports/optional/{stem}_optional.xlsx)',''])
    for q in qs:
        tier='Core' if q['core'] else 'Optional'
        choicehtml='<ol type="A">'+''.join('<li>'+H(c)+'</li>' for c in q['choices'])+'</ol>' if q['choices'] else '<p class="open-prompt">Open response · one or two sentences</p>'
        answer=(chr(64+q['correct_index'])+'. '+q['correct']) if q['correct_index'] else 'Suggested rubric (no automatic answer key)'
        placement=f'{q["when"].capitalize()} slide {q["slide_number"]}: {q["slide_title"]}'
        block.append(f'''<article class="question" id="{q['id']}" data-lecture="{q['lecture']}" data-tier="{tier.lower()}"><div class="qmeta"><span>{H(q['id'])}</span><span class="badge">{tier}</span><span>{H(q['purpose'])}</span><span>{q['seconds']} s including debrief</span></div><p class="placement"><a href="{q['slide_href']}">{H(placement)}</a> · <a href="{q['handout_href']}">Handout section</a></p><h3>{H(q['stem'])}</h3>{choicehtml}<div class="answer"><strong>{H(answer)}</strong><p>{H(q['explanation'])}</p></div><p><strong>Misconception to diagnose:</strong> {H(q['misconception'])}</p>{'<p><strong>Facilitation:</strong> '+H(q['activity'])+'</p>' if q['activity'] else ''}</article>''')
        md.extend([f'### {q["id"]} · {tier} · {q["purpose"]}','',f'**Placement:** [{placement}]({q["slide_href"]}) · [Handout]({q["handout_href"]})  ',f'**Time:** {q["seconds"]} seconds, including debrief.','',q['stem'],''])
        md.extend([f'- **{chr(65+i)}.** {c}' for i,c in enumerate(q['choices'])]);md.extend(['',f'**Answer:** {answer}','',q['explanation'],'',f'**Misconception:** {q["misconception"]}',''])
        if q['activity']:md.extend([f'**Facilitation:** {q["activity"]}',''])
    block.append('</section>');content.append('\n'.join(block))

css='''
:root{--ink:#233348;--muted:#607186;--accent:#007c70;--paper:#fffefa;--line:#d7e2e5}*{box-sizing:border-box}html{scroll-behavior:smooth;scroll-padding-top:90px}body{margin:0;background:var(--paper);color:var(--ink);font:17px/1.65 Georgia,serif}a{color:#006f79;text-underline-offset:3px}header{position:sticky;top:0;z-index:2;background:#102030;color:#fff;padding:12px max(20px,calc((100vw - 1100px)/2));display:flex;gap:20px;align-items:center;justify-content:space-between;font:14px/1.5 'Segoe UI',sans-serif}header a{color:#8de0d2}main{max-width:1100px;margin:45px auto 90px;padding:0 24px}h1,h2,h3,.eyebrow,.qmeta,.placement,table,label,button,select,.lecture-meta{font-family:'Segoe UI',Arial,sans-serif}h1{font-size:44px;line-height:1.15;letter-spacing:-.03em;margin:10px 0 25px}h2{font-size:30px;line-height:1.3;margin:0 0 10px}h3{font-size:22px;line-height:1.5;margin:18px 0}p{margin:12px 0}.eyebrow{color:var(--accent);font-size:12px;font-weight:700;letter-spacing:.14em;text-transform:uppercase}.intro{max-width:85ch}.instructions{background:#eff6f4;padding:22px 26px;border-radius:10px;margin:26px 0}.instructions h2{font-size:21px}.instructions p{font-size:16px}.table-wrap{overflow-x:auto}table{width:100%;border-collapse:collapse;font-size:14px;margin:20px 0 30px}th,td{padding:12px 10px;text-align:left;border-bottom:1px solid var(--line);vertical-align:top}th{background:#eaf2f2}th:first-child{width:38%}.filters{display:flex;gap:18px;flex-wrap:wrap;margin:25px 0;font-size:14px}.filters label{display:flex;gap:8px;align-items:center;max-width:100%;min-width:0}.filters select{min-width:0;max-width:100%}select,button{padding:8px 12px;border:1px solid #a7babf;background:white;border-radius:6px;color:var(--ink);font-size:14px}button{cursor:pointer}a:focus-visible,button:focus-visible,select:focus-visible{outline:3px solid #d78b25;outline-offset:3px}.lecture{padding-top:35px;margin-top:35px;border-top:3px solid var(--ink)}.lecture-meta,.placement{font-size:14px;color:var(--muted)}.question{padding:28px 0;border-bottom:1px solid var(--line);scroll-margin-top:90px}.qmeta{display:flex;align-items:center;flex-wrap:wrap;gap:12px;font-size:12px;color:var(--muted)}.qmeta>span:first-child{font-weight:700;color:var(--ink)}.badge{padding:3px 9px;border-radius:20px;background:#e4f2ee;color:#006c60}.question[data-tier=optional] .badge{background:#f7ecd8;color:#885817}ol{padding-left:28px}li{margin:8px 0;padding-left:4px}.answer{border-left:4px solid var(--accent);background:#eef7f3;padding:16px 20px;margin:23px 0}.answer p:last-child{margin-bottom:0}.open-prompt{font-style:italic;color:var(--muted)}.format-note{font-size:15px;color:var(--muted);margin:25px 0}.hidden{display:none!important}@media(max-width:650px){h1{font-size:33px}h2{font-size:26px}h3{font-size:20px}main{padding:0 18px;margin-top:28px}header{gap:10px;flex-wrap:wrap}.instructions{padding:18px}th,td{min-width:80px}th:first-child{min-width:210px}.qmeta{gap:8px}.answer{padding:14px}.placement{line-height:1.8}}@media print{@page{size:A4;margin:17mm}html{scroll-behavior:auto}body{background:white;font-size:10pt;line-height:1.45}header,.filters,.print-button{display:none!important}main{max-width:none;margin:0;padding:0}h1{font-size:24pt}h2{font-size:18pt}h3{font-size:12pt}.lecture{break-before:page;margin:0;padding-top:12px}.question{break-inside:avoid;padding:15px 0}.answer{padding:9px 13px;background:#f1f5f3;print-color-adjust:exact}.qmeta,.placement,.lecture-meta{font-size:8pt}a{color:inherit}table{font-size:8pt}th,td{padding:6px}.instructions{background:transparent;padding:0}.instructions p,.format-note{font-size:9pt}ol li{margin:4px 0}.question p{margin:7px 0}}
'''
script='''const lecture=document.querySelector('#lecture-filter'),tier=document.querySelector('#tier-filter');function filter(){document.querySelectorAll('.question').forEach(q=>q.classList.toggle('hidden',(lecture.value!=='all'&&q.dataset.lecture!==lecture.value)||(tier.value!=='all'&&q.dataset.tier!==tier.value)));document.querySelectorAll('.lecture').forEach(s=>s.classList.toggle('hidden',!s.querySelector('.question:not(.hidden)')));}lecture.addEventListener('change',filter);tier.addEventListener('change',filter);document.querySelector('[data-print]').addEventListener('click',()=>print());document.querySelectorAll('a[href^="#lecture-"]').forEach(a=>a.addEventListener('click',()=>{lecture.value='all';tier.value='all';filter();}));'''
doc=f'''<!doctype html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>AIML · Wooclap instructor guide</title><style>{css}</style></head><body><header><a href="../index.html">Course home</a><span>Instructor guide · answer keys included</span><button data-print>Print / save PDF</button></header><main><div class="eyebrow">AIML 2025/26 · Teaching companion</div><h1>Questions that make<br>the reasoning visible</h1><p class="intro">{intro}</p><div class="instructions"><h2>Use the questions within the lesson</h2><p>{method}</p><h2>Import into Wooclap</h2><p>{importnote}</p></div><div class="table-wrap"><table><thead><tr><th>Lecture</th><th>Core</th><th>Optional</th><th>Core budget</th><th>Import files</th></tr></thead><tbody>{''.join(rows)}</tbody></table></div><p class="format-note">{formatnote} See <a href="{fmt['official_documentation']}">official import instructions</a> and <a href="{fmt['template_copy']}">the reference template copy</a>. Checked 14 September 2026.</p><div class="filters"><label>Lecture <select id="lecture-filter"><option value="all">All lectures</option>{''.join(f'<option value="{l["number"]}">{l["number"]:02d} · {H(l["title"])}</option>' for l in bank['lectures'])}</select></label><label>Question set <select id="tier-filter"><option value="all">Core and optional</option><option value="core">Core only</option><option value="optional">Optional only</option></select></label></div>{''.join(content)}</main><script>{script}</script></body></html>'''
(HERE/'instructor-guide.html').write_text(doc,encoding='utf-8')
preface=['# AIML · Wooclap instructor guide','',intro,'',method,'','## Import instructions','',importnote,'',formatnote,'',f'[Official import documentation]({fmt["official_documentation"]}) · [Reference template]({fmt["template_copy"]})','']
(HERE/'instructor-guide.md').write_text('\n'.join(preface+md),encoding='utf-8')
report={'date':bank.get('revised',bank['created']),'questions':len(questions),'core':68,'optional':24,'types':dict(collections.Counter(q['type'] for q in questions)),'correct_positions':dict(collections.Counter(q['correct_index'] for q in questions if q['correct_index'])),'workbooks':32,'schema':fmt['columns'],'workbook_roundtrip':'passed','duplicate_stems':0,'slide_links':'all mapped to existing source slides and handout chapters','core_teaching_order':'passed','sources':sources,'live_wooclap_import_tested':True,'audience':bank.get('audience'),'live_revision_verification':'See live-revision-validation.json'}
(HERE/'validation.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
(HERE/'placement.json').write_text(json.dumps([{k:q[k] for k in ['id','pack','slide','slide_number','slide_title','when','slide_href','handout_href','core','seconds']} for q in questions],ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='sources'}))


# Keep the displayed slide prompts synchronized whenever the bank is rebuilt.
embed_path = HERE / 'embed-codes.json'
if embed_path.exists():
    embeds = json.loads(embed_path.read_text(encoding='utf-8'))
    mapping = {q['id']: (event, q) for event in embeds['events'] for q in event['questions']}
    assert set(mapping) == {q['id'] for q in questions}
    presentation = {'lectures': []}
    for lecture in bank['lectures']:
        selected = [q for q in questions if q['pack'] == lecture['pack']]
        presentation['lectures'].append({
            'pack': lecture['pack'], 'eventCode': mapping[selected[0]['id']][0]['eventCode'],
            'questions': [{**{k: q[k] for k in ['id','core','type','slide','when','purpose','stem','seconds','choices']},
                           'src': mapping[q['id']][1]['src']} for q in selected]
        })
    (HERE / 'slides-data.js').write_text('// Introductory undergraduate question prompts; no answer keys.\nwindow.AIML_WOOCLAP = ' + json.dumps(presentation, ensure_ascii=False, indent=2) + ';\n', encoding='utf-8')
